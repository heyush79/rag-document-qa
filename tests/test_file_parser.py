from __future__ import annotations

import io

import pytest

from app.core.exceptions import UnsupportedFileError
from app.services.file_parser import FileParser

HTML = b"""<!doctype html>
<html><head><title>Handbook</title>
<style>p { color: red }</style>
<script>alert("ignored")</script>
</head>
<body>
<h1>Leave policy</h1>
<p>Employees receive <b>24</b> paid leave days.</p>
<ul><li>Submit requests early.</li><li>Emergencies may be retroactive.</li></ul>
<table><tr><th>Team</th><th>Days</th></tr><tr><td>Platform</td><td>3</td></tr></table>
<noscript>needs javascript</noscript>
</body></html>
"""

CSV = b"name,role,location\nAda,Engineer,London\nGrace,Admiral,Arlington\n"


def docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Employees receive 24 paid leave days per year.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Team"
    table.rows[0].cells[1].text = "Days"
    table.rows[1].cells[0].text = "Platform"
    table.rows[1].cells[1].text = "3"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_html_keeps_visible_text_and_drops_scripts_and_styles() -> None:
    text = FileParser().parse("handbook.html", HTML)
    assert text.splitlines() == [
        "Handbook",
        "Leave policy",
        "Employees receive 24 paid leave days.",
        "Submit requests early.",
        "Emergencies may be retroactive.",
        "Team Days",
        "Platform 3",
    ]
    assert "alert" not in text and "color" not in text and "javascript" not in text


def test_htm_extension_is_html_too() -> None:
    assert "Leave policy" in FileParser().parse("page.htm", HTML)


def test_csv_becomes_one_block_per_row() -> None:
    text = FileParser().parse("people.csv", CSV)
    assert text == (
        "name: Ada\nrole: Engineer\nlocation: London\n\n"
        "name: Grace\nrole: Admiral\nlocation: Arlington"
    )


def test_csv_extra_values_without_a_column_are_dropped() -> None:
    text = FileParser().parse("odd.csv", b"a,b\n1,2,3\n4\n")
    assert text == "a: 1\nb: 2\n\na: 4\nb: "


def test_docx_paragraphs_then_tables() -> None:
    text = FileParser().parse("handbook.docx", docx_bytes())
    lines = [line for line in text.splitlines() if line]
    assert lines == [
        "Leave policy",
        "Employees receive 24 paid leave days per year.",
        "Team | Days",
        "Platform | 3",
    ]


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("renamed.pdf", b"just text pretending", "does not look like a PDF"),
        ("renamed.docx", b"just text pretending", "does not look like a DOCX"),
        ("broken.docx", b"PK\x03\x04 but not a real archive", "DOCX could not be parsed"),
        ("empty.html", b"<html><script>only()</script></html>", "no extractable text"),
        ("latin.csv", "name\ncaf\xe9".encode("latin-1"), "UTF-8"),
        ("notes.xlsx", b"PK\x03\x04", "Unsupported file type '.xlsx'"),
    ],
)
def test_rejected_files(filename, content, message) -> None:
    with pytest.raises(UnsupportedFileError, match=message):
        FileParser().parse(filename, content)


def test_unsupported_message_lists_the_types() -> None:
    with pytest.raises(UnsupportedFileError, match="TXT, MD, PDF, DOCX, HTML or CSV"):
        FileParser().parse("archive.zip", b"PK")


@pytest.mark.parametrize(
    ("filename", "content", "media_type"),
    [
        ("handbook.html", HTML, "text/html"),
        ("people.csv", CSV, "text/csv"),
        ("handbook.docx", None, "application/octet-stream"),
    ],
)
def test_new_types_upload_and_answer(client, filename, content, media_type) -> None:
    content = docx_bytes() if content is None else content
    response = client.post(
        "/api/v1/documents/upload", files={"file": (filename, content, media_type)}
    )
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "ready"
    question = "Where is Grace?" if filename.endswith(".csv") else "How many paid leave days?"
    answer = client.post("/api/v1/query", json={"question": question}).json()
    assert answer["sources"][0]["document_title"] == filename
