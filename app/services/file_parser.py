from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from pathlib import Path

from app.core.exceptions import UnsupportedFileError

TEXT_EXTENSIONS = {".txt", ".md"}
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".html", ".htm", ".csv"}
SUPPORTED_LIST = "TXT, MD, PDF, DOCX, HTML or CSV"

# Leading bytes that confirm the type, so a renamed file gets a 415 rather than a parser error.
# A DOCX file is a zip archive.
MAGIC_BYTES = {".pdf": b"%PDF-", ".docx": b"PK\x03\x04"}

_SKIPPED_TAGS = {"script", "style", "noscript", "template"}
_BLOCK_TAGS = {
    "p", "div", "br", "li", "ul", "ol", "dl", "dt", "dd", "h1", "h2", "h3", "h4", "h5", "h6",
    "tr", "table", "thead", "tbody", "section", "article", "header", "footer", "main", "nav",
    "aside", "blockquote", "pre", "hr", "figure", "figcaption", "address",
}
_CELL_TAGS = {"td", "th"}


class _HTMLText(HTMLParser):
    """Collects visible text: block elements start new lines, table cells are space-separated,
    and anything inside script or style is dropped."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in _SKIPPED_TAGS:
            self._skipping += 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")
        elif tag in _CELL_TAGS:
            self.parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIPPED_TAGS:
            self._skipping = max(0, self._skipping - 1)
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.parts.append(data)

    def text(self) -> str:
        lines = (" ".join(line.split()) for line in "".join(self.parts).splitlines())
        return "\n".join(line for line in lines if line)


def _decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedFileError("Text, HTML and CSV files must be UTF-8 encoded") from exc


def _pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise UnsupportedFileError("The PDF could not be parsed") from exc


def _docx_text(content: bytes) -> str:
    """Paragraphs in order, then each table with one row per line."""
    from docx import Document

    try:
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.append("")
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise UnsupportedFileError("The DOCX could not be parsed") from exc
    return "\n".join(parts)


def _html_text(content: bytes) -> str:
    parser = _HTMLText()
    parser.feed(_decode_utf8(content))
    parser.close()
    return parser.text()


def _csv_text(content: bytes) -> str:
    """One "column: value" line per field and a blank line between rows, so a row's fields stay
    together when the text is chunked."""
    try:
        rows = list(csv.DictReader(io.StringIO(_decode_utf8(content))))
    except csv.Error as exc:
        raise UnsupportedFileError("The CSV could not be parsed") from exc
    blocks = []
    for row in rows:
        lines = [
            f"{column.strip()}: {(value or '').strip()}"
            for column, value in row.items()
            if column is not None  # values beyond the header row have no column
        ]
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


class FileParser:
    SUPPORTED_EXTENSIONS = SUPPORTED_EXTENSIONS

    def parse(self, filename: str, content: bytes) -> str:
        extension = Path(filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFileError(
                f"Unsupported file type '{extension or 'unknown'}'. Use {SUPPORTED_LIST}."
            )
        magic = MAGIC_BYTES.get(extension)
        if magic and not content.startswith(magic):
            raise UnsupportedFileError(f"The file does not look like a {extension[1:].upper()}")

        if extension in TEXT_EXTENSIONS:
            text = _decode_utf8(content)
        elif extension == ".pdf":
            text = _pdf_text(content)
        elif extension == ".docx":
            text = _docx_text(content)
        elif extension == ".csv":
            text = _csv_text(content)
        else:
            text = _html_text(content)

        if not text.strip():
            raise UnsupportedFileError("The uploaded document contains no extractable text")
        return text
